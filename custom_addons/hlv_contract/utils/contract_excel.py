# -*- coding: utf-8 -*-
"""Xuất hợp đồng ra .xlsx: mỗi mẫu một sheet, như file Excel cũ (sheet HDMB, HDNT).

Dòng tiêu đề là tên biến (Ten_Ben_Mua, Don_Gia_1...) giống form cũ, nên file xuất ra vẫn
dùng làm nguồn Mail Merge của Word được. Giá trị số/ngày giữ kiểu thật để Excel tính tiếp.

build_workbook(sheets) -> bytes.
    sheets: list {'title': tên sheet, 'columns': [(khoá, kiểu)], 'rows': [dict khoá -> giá trị]}
    kiểu: 'text' | 'date' (datetime.date) | 'money' | 'number'. Khoá thiếu -> ô trống.
Trường hợp biên: sheets rỗng -> một sheet trống tên "HopDong"; tên sheet trùng/quá 31 ký tự
được cắt và đánh số để Excel mở được.
"""

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
_FORMATS = {'date': 'dd/mm/yyyy', 'money': '#,##0', 'number': '#,##0.###'}
_INVALID_SHEET_CHARS = re.compile(r'[\\/*?:\[\]]')


def _sheet_title(title, used):
    base = _INVALID_SHEET_CHARS.sub('-', title or 'HopDong')[:28] or 'HopDong'
    name, index = base, 2
    while name.lower() in used:
        name = '%s %d' % (base, index)
        index += 1
    used.add(name.lower())
    return name


def _fill_sheet(ws, sheet):
    columns = sheet['columns']
    ws.append([key for key, _kind in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = _HEADER_FILL
    for row in sheet['rows']:
        ws.append([row.get(key) for key, _kind in columns])
    for idx, (_key, kind) in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=idx).column_letter
        ws.column_dimensions[letter].width = 18
        if kind in _FORMATS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = _FORMATS[kind]
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions


def build_workbook(sheets):
    """Xem docstring đầu module."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sheet in sheets or [{'title': 'HopDong', 'columns': [], 'rows': []}]:
        _fill_sheet(wb.create_sheet(_sheet_title(sheet['title'], used)), sheet)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
