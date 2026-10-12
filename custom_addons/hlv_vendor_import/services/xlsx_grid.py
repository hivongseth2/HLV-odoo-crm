# -*- coding: utf-8 -*-
"""Đọc file Excel NCC thành lưới giá trị. Util thuần: vào bytes, ra list — không đụng env."""

import datetime
import io

# File tồn kho / bảng giá NCC thường vài trăm tới vài nghìn dòng; chặn trên để file lỗi (cả triệu dòng
# trống có định dạng) không làm treo worker.
MAX_ROWS = 20000


class GridError(ValueError):
    """File không đọc được; thông điệp viết cho người dùng đọc."""


def _cell(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, datetime.datetime):
        return value.strftime("%d/%m/%Y")
    return value


def read_grid(data, sheet_name=None):
    """Đọc một sheet của file .xlsx.

    Nhận: ``data`` — bytes của file; ``sheet_name`` — tên sheet, rỗng = sheet đầu tiên.
    Trả: ``(tên các sheet, lưới)``; lưới là list các dòng, mỗi dòng là list giá trị (None cho ô trống,
    chuỗi đã strip, số giữ nguyên, ngày thành "dd/mm/YYYY"). Ô gộp (merge) được điền giá trị ô đầu vào
    mọi ô của vùng gộp: tiêu đề hai tầng gộp dọc / gộp ngang mới ghép được.
    Lỗi: ``GridError`` khi không phải .xlsx, không có sheet tên đó, hoặc vượt ``MAX_ROWS`` dòng.
    """
    import openpyxl  # Odoo đã cần openpyxl (base_import); import muộn để util nhẹ khi không dùng.
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    try:
        # Không dùng read_only: chế độ đó không cho biết vùng gộp.
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError):
        raise GridError("Không đọc được file. Chỉ nhận Excel .xlsx — file .xls / .csv thì mở bằng Excel rồi lưu lại dạng .xlsx.")
    names = workbook.sheetnames
    if sheet_name and sheet_name not in names:
        raise GridError("File không có sheet \"%s\" (có: %s)." % (sheet_name, ", ".join(names)))
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    if sheet.max_row > MAX_ROWS:
        raise GridError("Sheet \"%s\" có %s dòng, quá giới hạn %s." % (sheet.title, sheet.max_row, MAX_ROWS))
    grid = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    for merged in sheet.merged_cells.ranges:
        top = grid[merged.min_row - 1][merged.min_col - 1] if merged.min_row <= len(grid) else None
        for row in range(merged.min_row - 1, min(merged.max_row, len(grid))):
            for col in range(merged.min_col - 1, min(merged.max_col, len(grid[row]))):
                grid[row][col] = top
    return names, grid
